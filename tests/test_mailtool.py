from datetime import date
from email.message import EmailMessage

import pytest

from mailtool.attachments import safe_filename, save_attachments
from mailtool.config import Settings
from mailtool.imap_client import ReadOnlyIMAP4_SSL, ReadOnlyViolation, SearchCriteria, build_search
from mailtool.pdf_extract import looks_garbled_arabic


# --- search -------------------------------------------------------------

def test_search_last_two_days_includes_today_and_yesterday():
    tokens, literal = build_search(
        SearchCriteria(subject="Daily Morning Briefing Minutes", days=2), today=date(2026, 10, 22))
    assert tokens == ["SUBJECT", '"Daily Morning Briefing Minutes"', "SINCE", "21-Oct-2026"]
    assert literal is None


def test_search_sender_and_date_range():
    tokens, _ = build_search(SearchCriteria(sender="exec.secretary@sunnydaysegypt.com",
                                            since=date(2026, 10, 1), before=date(2026, 10, 5)))
    assert tokens == ["FROM", '"exec.secretary@sunnydaysegypt.com"',
                      "SINCE", "01-Oct-2026", "BEFORE", "05-Oct-2026"]


def test_search_quotes_are_escaped():
    tokens, _ = build_search(SearchCriteria(subject='say "hi"'))
    assert tokens == ["SUBJECT", '"say \\"hi\\""']


def test_search_arabic_subject_goes_in_literal_last():
    tokens, literal = build_search(SearchCriteria(subject="محضر", days=1), today=date(2026, 10, 22))
    assert tokens[-1] == "SUBJECT"
    assert literal == "محضر".encode()


def test_search_rejects_two_non_ascii_terms():
    with pytest.raises(ValueError):
        build_search(SearchCriteria(sender="سيف", subject="محضر"))


# --- read-only guard ----------------------------------------------------

class _Recorder(ReadOnlyIMAP4_SSL):
    """Skip the network; record what would have been sent."""

    def __init__(self):
        self.sent = []

    def _simple_command(self, name, *args):
        return self._command(name, *args)


@pytest.fixture
def conn(monkeypatch):
    c = _Recorder()
    monkeypatch.setattr("imaplib.IMAP4._command",
                        lambda self, name, *args: self.sent.append((name, args)) or "tag")
    return c


@pytest.mark.parametrize("name,args", [
    ("STORE", ("1", "+FLAGS", "(\\Seen)")),
    ("UID", ("STORE", "1", "+FLAGS", "(\\Deleted)")),
    ("UID", ("MOVE", "1", "Trash")),
    ("UID", ("COPY", "1", "Trash")),
    ("EXPUNGE", ()),
    ("DELETE", ("INBOX",)),
    ("SELECT", ("INBOX",)),
    ("UID", ("FETCH", "1", "(BODY[])")),
    ("UID", ("FETCH", "1", "(RFC822)")),
    ("FETCH", ("1", "(RFC822.TEXT)")),
])
def test_modifying_commands_are_blocked(conn, name, args):
    with pytest.raises(ReadOnlyViolation):
        conn._command(name, *args)
    assert conn.sent == []


@pytest.mark.parametrize("name,args", [
    ("UID", ("FETCH", "1", "(BODY.PEEK[])")),
    ("UID", ("FETCH", "1", "(RFC822.HEADER RFC822.SIZE)")),
    ("UID", ("SEARCH", "SUBJECT", '"BODY[ RFC822"')),
    ("EXAMINE", ("INBOX",)),
])
def test_read_commands_pass(conn, name, args):
    conn._command(name, *args)
    assert conn.sent == [(name, args)]


def test_select_read_write_is_refused(conn):
    with pytest.raises(ReadOnlyViolation):
        conn.select("INBOX", readonly=False)


# --- attachments --------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("../../etc/passwd", "passwd"),
    ("..\\..\\evil.pdf", "evil.pdf"),
    ("Daily Beriefing 22102026 .pdf", "Daily Beriefing 22102026.pdf"),
    ('a:b*c?"d<e>f|g.pdf', "a_b_c__d_e_f_g.pdf"),
    ("..", "attachment"),
    (None, "attachment"),
    ("تقرير.pdf", "تقرير.pdf"),
])
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def _msg_with(*attachments):
    msg = EmailMessage()
    msg["Subject"] = "t"
    msg.set_content("body")
    for name, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    return msg


def test_save_attachments_by_date_pdf_only(tmp_path):
    msg = _msg_with(("../min.pdf", b"%PDF-1.5 a", "application", "pdf"),
                    ("logo.png", b"\x89PNG", "image", "png"))
    saved = save_attachments(msg, date(2026, 10, 22), tmp_path, documents_only=True)
    assert [s.path for s in saved] == [tmp_path / "2026-10-22" / "min.pdf"]
    assert saved[0].path.read_bytes() == b"%PDF-1.5 a"


def test_save_attachments_dedupes_and_avoids_overwrite(tmp_path):
    d = date(2026, 10, 22)
    save_attachments(_msg_with(("m.pdf", b"one", "application", "pdf")), d, tmp_path)
    again = save_attachments(_msg_with(("m.pdf", b"one", "application", "pdf")), d, tmp_path)
    other = save_attachments(_msg_with(("m.pdf", b"two", "application", "pdf")), d, tmp_path)
    assert again[0].skipped_existing
    assert other[0].path.name == "m (2).pdf"
    assert (tmp_path / "2026-10-22" / "m.pdf").read_bytes() == b"one"


# --- secrets ------------------------------------------------------------

def test_settings_repr_hides_secrets():
    s = Settings(imap_user="u@example.com", imap_password="abcdefghijklmnop",
                 anthropic_api_key="sk-ant-secret")
    assert "abcdefghijklmnop" not in repr(s)
    assert "sk-ant-secret" not in repr(s)


# --- garbled Arabic detection -------------------------------------------

def test_good_arabic_is_not_garbled():
    text = "ضرورة توفير صندوق إسعافات أولية في جميع المطاعم والبارات مع التأكد من توافره على مدار اليوم " * 2
    assert not looks_garbled_arabic(text)


def test_broken_text_layer_is_garbled():
    text = "ضورة توفن صندوق إسعافات أولية يف جميع المطاعم والبارات مع إعطاء هذا األمر أهمية قصوى " * 2
    assert looks_garbled_arabic(text.replace("مع", "مغ"))


# --- Excel / Word -------------------------------------------------------

def _make_xlsx(path):
    import datetime as dt
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "22.10"
    ws.append(["Hotel", "Occupancy", "Date"])
    ws.append(["El Palacio", 99.75, dt.datetime(2026, 10, 22)])
    ws.append([None, None, None])
    ws.append(["Palma", 78, None])
    wb.create_sheet("23.10").append(["ميريت", "85%"])
    wb.save(path)


def test_extract_xlsx_all_sheets(tmp_path):
    from mailtool.office_extract import extract_office

    path = tmp_path / "report.xlsx"
    _make_xlsx(path)
    text = extract_office(path)
    assert "===== Sheet: 22.10 (3 rows) =====" in text
    assert "El Palacio\t99.75\t2026-10-22" in text
    assert "Palma\t78" in text.splitlines()
    assert "ميريت\t85%" in text


def test_xls_name_with_xlsx_content_is_read(tmp_path):
    from mailtool.office_extract import extract_office

    path = tmp_path / "Competition Analysis.xls"
    _make_xlsx(path)
    assert "El Palacio" in extract_office(path)


def test_extract_docx_keeps_tables_in_order(tmp_path):
    import docx
    from mailtool.office_extract import extract_office

    document = docx.Document()
    document.add_paragraph("Duty Manager Incident Log 22/10/2026")
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Time", "Room", "Incident"]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["23:40", "1204", "Guest\nlocked out"]):
        cell.text = value
    document.add_paragraph("Signed: Night Manager")
    path = tmp_path / "log.docx"
    document.save(path)

    text = extract_office(path)
    assert text.index("Incident Log") < text.index("[table]") < text.index("Signed")
    assert "23:40\t1204\tGuest locked out" in text


def test_save_attachments_keeps_office_documents(tmp_path):
    msg = _msg_with(
        ("sheet.xlsx", b"PK..", "application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("old.xls", b"\xd0\xcf..", "application", "vnd.ms-excel"),
        ("log.docx", b"PK..", "application", "vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("logo.png", b"\x89PNG", "image", "png"),
    )
    saved = save_attachments(msg, date(2026, 10, 23), tmp_path, documents_only=True)
    assert sorted(s.path.name for s in saved) == ["log.docx", "old.xls", "sheet.xlsx"]
