"""Extract text from Excel (.xlsx/.xlsm/.xls) and Word (.docx) files."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
LEGACY_EXCEL_SUFFIXES = {".xls"}
WORD_SUFFIXES = {".docx"}


def extract_office(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return _extract_xlsx(path)
    if suffix in LEGACY_EXCEL_SUFFIXES:
        return _extract_xls(path)
    if suffix in WORD_SUFFIXES:
        return _extract_docx(path)
    raise ValueError(f"unsupported file type: {path.name}")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.datetime, dt.date)):
        if isinstance(value, dt.datetime) and value.time() == dt.time(0):
            value = value.date()
        return value.isoformat()
    return " ".join(str(value).split())  # collapse newlines inside a cell


def _rows_to_text(rows) -> list[str]:
    lines = []
    for row in rows:
        cells = [_cell_text(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append("\t".join(cells))
    return lines


def _sheets_to_text(sheets: list[tuple[str, list[str]]]) -> str:
    return "\n\n".join(f"===== Sheet: {name} ({len(lines)} rows) =====\n" + "\n".join(lines)
                       for name, lines in sheets)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    # A file handle, not the path: openpyxl rejects names that don't end in .xlsx.
    # data_only: use the values Excel last calculated, not the formulas.
    with path.open("rb") as fh:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        try:
            sheets = [(ws.title, _rows_to_text(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
        finally:
            wb.close()
    return _sheets_to_text(sheets)


def _extract_xls(path: Path) -> str:
    import xlrd

    if path.read_bytes()[:2] == b"PK":  # an .xlsx saved with an .xls name
        return _extract_xlsx(path)
    book = xlrd.open_workbook(path)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row = []
            for cell in sheet.row(r):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value if cell.value != "" else None)
            rows.append(row)
        sheets.append((sheet.name, _rows_to_text(rows)))
    return _sheets_to_text(sheets)


def _extract_docx(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    parts: list[str] = []
    # Walk the body in order so tables stay next to the paragraphs around them.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    text = _cell_text(cell.text)
                    if not cells or cells[-1] != text:  # merged cells repeat their text
                        cells.append(text)
                rows.append(cells)
            lines = _rows_to_text(rows)
            if lines:
                parts.append("[table]\n" + "\n".join(lines) + "\n[/table]")
    images = sum(1 for rel in document.part.rels.values() if "image" in rel.reltype)
    if images:
        parts.append(f"[{images} embedded image(s) not extracted]")
    return "\n\n".join(parts)
